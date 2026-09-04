#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""parse_xiugai_moban.py — 解析用户改完拖回的「审核修改模板.xlsx」，对隐藏「基线」sheet 做 diff

入参: xlsx 路径 [--baseline baselineId 可选（服务端已去掉，默认读隐藏「基线」sheet）]
出参(JSON): {"remove":[], "add":[], "assignChanges":[{"clauseId","from","to"}], "invalid":[], "empty":[]}
"""
import sys, json, argparse
import sys
try:
    # Windows 下强制 UTF-8，避免 stdin/stdout 按 GBK 解码导致中文乱码
    sys.stdin.reconfigure(encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from openpyxl import load_workbook

KEY = "条款实例ID"


def _s(v):
    from datetime import datetime, date
    return v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else v

def load_rows(ws, start_row=2, layout="main"):
    """layout=main: [区域,项目,子要素,条款名称,分派人,条款实例ID]；layout=baseline: [clauseId,assignee,region,project,subElement,clauseName]"""
    rows = []
    for r in ws.iter_rows(min_row=start_row, values_only=True):
        vals = [_s(c) for c in (list(r) + [None] * 6)]
        if layout == "baseline":
            cid, assignee, region, project, sub, cname = vals[:6]
        else:
            region, project, sub, cname, assignee, cid = vals[:6]
            if isinstance(cid, str) and cid.startswith("="):
                cid = ""
        if not any([region, project, sub, cname, assignee, cid]):
            continue
        rows.append({"clauseId": cid or "", "region": region or "", "project": project or "",
                     "subElement": sub or "", "clauseName": cname or "", "assignee": assignee or ""})
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--baseline", default=None)
    a = ap.parse_args()

    wb = load_workbook(a.path, data_only=False)
    current = load_rows(wb["审核修改模板"])
    baseline = load_rows(wb["基线"], layout="baseline") if "基线" in wb.sheetnames else []

    # F 为公式时无缓存值：从隐藏「级联字典」按 区域|项目|子要素|条款名称 四元组反查条款ID
    four_map = {}
    if "级联字典" in wb.sheetnames:
        for rr in wb["级联字典"].iter_rows(min_row=2, values_only=True):
            vals = list(rr) + [None] * 9
            if vals[7] and vals[8]:
                four_map[vals[7]] = vals[8]
    for x in current:
        if not x["clauseId"]:
            x["clauseId"] = four_map.get("|".join([x["region"], x["project"], x["subElement"], x["clauseName"]]), "")

    base_by_id = {x["clauseId"]: x for x in baseline if x["clauseId"]}
    cur_by_id = {x["clauseId"]: x for x in current if x["clauseId"]}

    remove = [x["clauseId"] for x in baseline if x["clauseId"] and x["clauseId"] not in cur_by_id]
    add = [x["clauseId"] for x in current if x["clauseId"] and x["clauseId"] not in base_by_id]
    assign_changes = []
    for cid, cur in cur_by_id.items():
        b = base_by_id.get(cid)
        if b and b["assignee"] != cur["assignee"]:
            assign_changes.append({"clauseId": cid, "from": b["assignee"], "to": cur["assignee"]})
    invalid = [x["clauseId"] for x in current if not x["clauseId"]]
    empty = [x["clauseId"] for x in current if x["clauseId"] and not x["assignee"]]

    print(json.dumps({
        "remove": remove, "add": add, "assignChanges": assign_changes,
        "invalid": invalid, "empty": empty,
        "summary": f"删 {len(remove)} 条；增 {len(add)} 条；分派变更 {len(assign_changes)} 条；无ID {len(invalid)}；分派留空 {len(empty)}",
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

