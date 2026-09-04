#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""parse_fanwei_peizhi.py — 解析用户填回的「审核范围配置.xlsx」

入参: xlsx 路径 [--tree clauseTree.json 可选]
出参(JSON): {
  "period": {"granularity": "", "start": "", "end": ""},
  "factory": "",
  "clauses": [{"clauseId":"","clauseName":"","region":"","project":"","subElement":"","assignee":""}],
  "issues": ["行N: 描述"]
}
"""
import sys, json, argparse
import sys
try:
    # Windows 下强制 UTF-8，避免 stdin/stdout 按 GBK 解码导致中文乱码
    sys.stdin.reconfigure(encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from openpyxl import load_workbook

TASK_HEADER = ["周期粒度", "周期起", "周期止", "制造基地"]
DETAIL_HEADER = ["区域", "项目", "子要素", "条款名称", "分派人", "条款实例ID"]


def build_index(clause_tree):
    """4 元组(区域,项目,子要素,条款名) -> clauseId"""
    idx = {}
    for r in clause_tree:
        for p in r.get("projects", []):
            for s in p.get("subElements", []):
                for c in s.get("clauses", []):
                    idx[(r["region"], p["project"], s["subElement"], c["clauseName"])] = c["clauseId"]
    return idx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--tree", default=None, help="条款子树 JSON 文件")
    a = ap.parse_args()

    wb = load_workbook(a.path, data_only=False)
    ws = wb["审核范围配置"]
    def _s(v):
        from datetime import datetime, date
        return v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else v
    rows = [[_s(c) for c in r] for r in ws.iter_rows(values_only=True)]

    # 任务级：行1 表头，行2 值
    period = {"granularity": rows[1][0] or "", "start": rows[1][1] or "", "end": rows[1][2] or ""}
    factory = rows[1][3] or ""

    idx = {}
    if a.tree:
        with open(a.tree, "r", encoding="utf-8") as f:
            idx = build_index(json.load(f))
    four_map = {}
    if "级联字典" in wb.sheetnames:
        for rr in wb["级联字典"].iter_rows(min_row=2, values_only=True):
            vals = list(rr) + [None] * 9
            if vals[7] and vals[8]:
                four_map[vals[7]] = vals[8]

    clauses, issues = [], []
    seen = {}
    task_ok = bool(period["granularity"] and period["start"] and period["end"] and factory)
    if not task_ok:
        issues.append("任务级四格（周期粒度/起/止/基地）未填全")

    for i, r in enumerate(rows[4:], start=5):
        region, project, sub, cname, assignee = (list(r) + [None] * 5)[:5]
        cid = (list(r) + [None] * 6)[5] or ""
        if isinstance(cid, str) and cid.startswith("="):
            cid = ""
        if not any([region, project, sub, cname, assignee, cid]):
            continue
        row_no = i
        if not task_ok:
            issues.append(f"行{row_no}: 任务级四格（周期粒度/起/止/基地）未填全")
        if not assignee:
            issues.append(f"行{row_no}: 分派人未填")
        resolved = cid or four_map.get("|".join([region, project, sub, cname]), "")
        if not resolved and idx:
            resolved = idx.get((region, project, sub, cname), "")
        if not resolved and (idx or four_map):
            issues.append(f"行{row_no}: 区域/项目/子要素/条款 四元组不合法或未选齐")
        if resolved:
            if resolved in seen:
                issues.append(f"行{row_no}: 条款 {resolved} 重复")
            seen[resolved] = True
        clauses.append({
            "clauseId": resolved, "clauseName": cname or "",
            "region": region or "", "project": project or "",
            "subElement": sub or "", "assignee": assignee or "",
        })
    print(json.dumps({"period": period, "factory": factory, "clauses": clauses, "issues": issues}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()