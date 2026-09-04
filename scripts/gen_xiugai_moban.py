#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""gen_xiugai_moban.py — 生成「审核修改模板.xlsx」（用户点 ④改条款 / ⑤改分派人时）

入参(JSON, argv[1] 或 stdin): {
  "clauses": [ {"clauseId":"LJ-01","clauseName":"问题传递及时性","region":"零部件","project":"质量数据采集","subElement":"质量信息传递","assigneeId":3,"assigneeName":"王五","nodeType":"leaf"} ],
  "clauseTree": [...],
  "out": "审核修改模板.xlsx"
}
出参(JSON): {"file":..., "baselineId": "md5(原始条款快照)"}
基线快照写入隐藏「基线」sheet（服务端 /baselines 已去掉，客户端兜底 diff）。
"""
import sys, json, hashlib
import sys
try:
    # Windows 下强制 UTF-8，避免 stdin/stdout 按 GBK 解码导致中文乱码
    sys.stdin.reconfigure(encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.workbook.defined_name import DefinedName

HEADER = ["区域", "项目", "子要素", "条款名称", "分派人", "条款实例ID"]


def safe_key(s):
    """生成 Excel 命名区域安全内部键（确定性，跨进程稳定）"""
    return "k" + hashlib.md5(s.encode("utf-8")).hexdigest()[:10]


def build_dict(clause_tree):
    regions, region_keys, project_map, sub_map, clause_map = [], {}, {}, {}, {}
    for r in clause_tree:
        regions.append(r["region"])
        rk = safe_key(r["region"])
        region_keys[r["region"]] = rk
        project_map[rk] = []
        for p in r.get("projects", []):
            pk = safe_key(r["region"] + p["project"])
            project_map[rk].append((p["project"], pk))
            sub_map[pk] = []
            for s in p.get("subElements", []):
                sk = safe_key(r["region"] + p["project"] + s["subElement"])
                sub_map[pk].append((s["subElement"], sk))
                clause_map[sk] = [(c["clauseId"], c["clauseName"]) for c in s.get("clauses", [])]
    return regions, region_keys, project_map, sub_map, clause_map


def main():
    raw = sys.argv[1] if len(sys.argv) > 1 else sys.stdin.read()
    data = json.loads(raw) if raw.strip() else {}
    clauses = data.get("clauses", [])
    clause_tree = data.get("clauseTree", [])
    out = data.get("out") or "审核修改模板.xlsx"

    baseline = [{"clauseId": c.get("clauseId"), "assignee": c.get("assigneeName", ""),
                 "region": c.get("region"), "project": c.get("project"),
                 "subElement": c.get("subElement"), "clauseName": c.get("clauseName")} for c in clauses]
    baseline_id = hashlib.md5(json.dumps(baseline, ensure_ascii=False).encode("utf-8")).hexdigest()[:16]

    regions, region_keys, project_map, sub_map, clause_map = build_dict(clause_tree)

    wb = Workbook()
    ws = wb.active
    ws.title = "审核修改模板"
    ws.append(HEADER)
    for c in clauses:
        ws.append([c.get("region", ""), c.get("project", ""), c.get("subElement", ""),
                   c.get("clauseName", ""), c.get("assigneeName", ""), c.get("clauseId", "")])
    # 3 行空追加行
    for _ in range(3):
        ws.append(["", "", "", "", "", ""])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center")
    for col in ("G", "H", "I", "J"):
        ws.column_dimensions[col].hidden = True

    # ---- 级联字典（隐藏）----
    d = wb.create_sheet("级联字典")
    d.sheet_state = "hidden"
    d.append(["KEY", "区域", "项目", "子要素", "条款ID", "条款名", "条款名键", "四元组", "条款ID4", "区域|项目", "项目|子要素"])
    region_key_row, project_key_row, sub_key_row = {}, {}, {}
    r = 2
    for region in regions:
        rk = region_keys[region]
        region_key_row[rk] = r
        d.cell(r, 1, rk); d.cell(r, 2, region); r += 1
        for project, pk in project_map[rk]:
            project_key_row[pk] = r
            d.cell(r, 1, pk); d.cell(r, 3, project); d.cell(r, 10, rk + "|" + project); r += 1
            for sub, sk in sub_map[pk]:
                sub_key_row[sk] = r
                d.cell(r, 1, sk); d.cell(r, 4, sub); d.cell(r, 11, pk + "|" + sub); r += 1
                for cid, cname in clause_map[sk]:
                    d.cell(r, 1, "c_" + cid); d.cell(r, 5, cid); d.cell(r, 6, cname)
                    d.cell(r, 7, sk + "|" + cname)
                    d.cell(r, 8, f"{region}|{project}|{sub}|{cname}"); d.cell(r, 9, cid); r += 1
    last = r - 1
    if regions:
        wb.defined_names["regions"] = DefinedName("regions", attr_text=f"级联字典!$B$2:$B${region_key_row[region_keys[regions[-1]]]}")
    for rk, row in region_key_row.items():
        ps = project_map[rk]
        if ps:
            wb.defined_names["projects_" + rk] = DefinedName("projects_" + rk, attr_text=f"级联字典!$C${project_key_row[ps[0][1]]}:$C${project_key_row[ps[-1][1]]}")
    for pk, row in project_key_row.items():
        ss = sub_map[pk]
        if ss:
            wb.defined_names["subs_" + pk] = DefinedName("subs_" + pk, attr_text=f"级联字典!$D${sub_key_row[ss[0][1]]}:$D${sub_key_row[ss[-1][1]]}")
    for sk, row in sub_key_row.items():
        cs = clause_map[sk]
        if cs:
            wb.defined_names["clauses_" + sk] = DefinedName("clauses_" + sk, attr_text=f"级联字典!$F${row}:$F${row + len(cs) - 1}")

    # 明细行：A~D 已有行只读锁定、追加行级联下拉；E 可改；
    # F（条款实例ID）取消隐藏、整列锁定只读，用「区域|项目|子要素|条款名称」四元组 VLOOKUP 自动带出
    for row in range(2, ws.max_row + 1):
        A, B, C, D, E, F = (f"{c}{row}" for c in "ABCDEF")
        G, H, I = f"G{row}", f"H{row}", f"I{row}"
        ws[G] = f'=IF({A}="","",INDEX(级联字典!$A$2:$A${last},MATCH({A},级联字典!$B$2:$B${last},0)))'
        ws[H] = f'=IF(OR({G}="",{B}=""),"",INDEX(级联字典!$A$2:$A${last},MATCH({G}&"|"&{B},级联字典!$J$2:$J${last},0)))'
        ws[I] = f'=IF(OR({H}="",{C}=""),"",INDEX(级联字典!$A$2:$A${last},MATCH({H}&"|"&{C},级联字典!$K$2:$K${last},0)))'
        ws[F] = f'=IF({D}="","",IFERROR(VLOOKUP({A}&"|"&{B}&"|"&{C}&"|"&{D},级联字典!$H$2:$I${last},2,0),""))'
        ws[F].protection = Protection(locked=True)
        is_append = row > len(clauses) + 1
        if is_append:
            dvA = DataValidation(type="list", formula1="=regions", allow_blank=True, showDropDown=False)
            dvB = DataValidation(type="list", formula1=f'=INDIRECT("projects_"&{G})', allow_blank=True, showDropDown=False)
            dvC = DataValidation(type="list", formula1=f'=INDIRECT("subs_"&{H})', allow_blank=True, showDropDown=False)
            dvD = DataValidation(type="list", formula1=f'=INDIRECT("clauses_"&{I})', allow_blank=True, showDropDown=False)
            for dv, col in ((dvA, A), (dvB, B), (dvC, C), (dvD, D)):
                ws.add_data_validation(dv)
                dv.add(ws[col])
            for col in (A, B, C, D):
                ws[col].protection = Protection(locked=False)
        else:
            for col in (A, B, C, D):
                ws[col].protection = Protection(locked=True)
        ws[E].protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.protection.password = "wb-audit"
    ws.protection.enable()

    # ---- 隐藏「基线」sheet（diff 基准）----
    b = wb.create_sheet("基线")
    b.sheet_state = "hidden"
    b.append(["clauseId", "assignee", "region", "project", "subElement", "clauseName"])
    for x in baseline:
        b.append([x["clauseId"], x["assignee"], x["region"], x["project"], x["subElement"], x["clauseName"]])

    wb.save(out)
    print(json.dumps({"file": out, "baselineId": baseline_id, "clauseCount": len(clauses)}, ensure_ascii=False))


if __name__ == "__main__":
    main()